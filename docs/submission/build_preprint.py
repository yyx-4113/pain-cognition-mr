#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build the Research Square PREPRINT package (version 1) from manuscript_draft.md.

This is a PREPRINT, not a journal submission. The CHARLS [PENDING DATA] sections
are kept verbatim and explicitly framed as planned/prospective. A preprint banner
and a version note are injected on the title page.

Outputs into ./_upload_preprint :
    Preprint_Manuscript.docx        title page + preprint banner + abstract + IMRaD
                                     + references + declarations + figure legends
    Preprint_Metadata.md            every field for the Research Square upload form
    Preprint_Guide.md               (中文) step-by-step upload + versioning guide
    Supplementary_Information.docx  (copied from _upload)
    STROBE-MR_checklist.docx        (copied from _upload)
    figures/Fig1..Fig5.{pdf,png}    (copied from _upload)

Reuses the document engine from build_srep_pack.py (no value is recomputed).
"""
import os, re, shutil
from docx.enum.text import WD_ALIGN_PARAGRAPH
import build_srep_pack as B

HERE = os.path.dirname(os.path.abspath(__file__))
OUT  = os.path.join(HERE, "_upload_preprint")
OUTF = os.path.join(OUT, "figures")
SRC_UP = os.path.join(HERE, "_upload")
os.makedirs(OUT, exist_ok=True)
os.makedirs(OUTF, exist_ok=True)

PREPRINT_NOTE = (
    "Preprint note. This is a preprint of a manuscript that has not undergone peer "
    "review; it is posted to solicit feedback. The China Health and Retirement "
    "Longitudinal Study (CHARLS) longitudinal validation component is ongoing and is "
    "marked [PENDING DATA] throughout: those sections describe planned analyses, not "
    "results. A revised version (version 2) will be posted when CHARLS data access is "
    "granted and incorporated, after which the manuscript will be submitted to a "
    "peer-reviewed journal (Scientific Reports, Nature Portfolio)."
)

def build_preprint_manuscript(text):
    doc = B.new_document()
    # title page + author + abstract (everything before "## 1. Introduction")
    doc.add_paragraph()  # small top spacer
    B.emit_markdown(doc, text[:text.index("## 1. Introduction")])
    # preprint banner
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("PREPRINT · VERSION 1 · NOT PEER-REVIEWED")
    r.bold = True; r.font.size = B.Pt(11) if hasattr(B, "Pt") else None
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("Posted on Research Square — a DOI will be assigned on posting.")
    r2.italic = True
    # version note
    pn = doc.add_paragraph()
    pn.paragraph_format.space_before = B.Pt(6)
    B.add_runs(pn, PREPRINT_NOTE)
    doc.add_page_break()
    # IMRaD body
    intro   = text.index("## 1. Introduction")
    strobe  = text.index("## 6. STROBE-MR compliance")
    refs    = text.index("## References")
    audit   = text.index("## 7. Audit trail")
    legends = text.index("## 8. Figure captions")
    B.emit_markdown(doc, text[intro:strobe])
    B.emit_markdown(doc, text[refs:audit])
    for label, txt in B.DECLARATIONS:
        pp = doc.add_paragraph()
        rr = pp.add_run(label + ". "); rr.bold = True
        B.add_runs(pp, txt)
        pp.paragraph_format.space_before = B.Pt(6)
    B.emit_markdown(doc, "## Reporting guidelines\n\n" + B.REPORTING_NOTE)
    B.emit_markdown(doc, text[legends:])
    path = os.path.join(OUT, "Preprint_Manuscript.docx")
    doc.save(path)
    return path

METADATA = """# Research Square — Preprint Upload Metadata (Version 1)

> Fill the Research Square submission form with the fields below. Everything is
> taken verbatim from the manuscript; no value is invented.

## Manuscript type
Preprint (original research article)

## Title
Genetically predicted multisite chronic pain is associated with lower cognitive performance: a bidirectional Mendelian randomization study with NHANES contextual evidence

## Abstract
Chronic pain and cognitive impairment are rising ageing-related burdens. Observational links between multisite chronic pain (MCP) and cognitive decline are confounded by reverse causation and comorbidity. We used bidirectional two-sample MR to assess the association of genetically predicted MCP with cognitive outcomes. MCP instruments (51 SNPs; P<5x10-8) came from a UK Biobank GWAS (n=387,649). Outcomes were cognitive performance (UKB-derived, n=257,841), Alzheimer's disease (IGAP ieu-a-297, n=54,162) and all-cause dementia (FinnGen, n=216,771). Genetically predicted MCP was associated with lower cognitive performance (IVW beta=-0.361, 95% CI -0.424 to -0.299, p=3.65x10-30) with high heterogeneity (I2~83.5%) and no directional pleiotropy. The Steiger test was non-significant (p=0.23) and reverse MR significant (beta=-0.171, p=8.85x10-58), so direction is uncertain. A non-UKB overlap-robustness re-interpretation (IGAP AD: beta=0.466, p=0.008) corroborated the association independently of UK Biobank sample overlap. No mediator was significant. Genetic evidence is compatible with an MCP-lower-cognition association, but the direction is unresolved and the cognitive result is limited by UKB sample overlap.

## Authors
1. Yongxin Yang, B.M. (single author; corresponding author)
   - ORCID: 0009-0004-9698-6552
   - Affiliation: The Second Affiliated Hospital of Fujian University of Traditional Chinese Medicine, Fuzhou, Fujian 350003, China
   - Email: 960856791@qq.com

## Keywords (5-6)
Mendelian randomization; chronic pain; multisite chronic pain; cognitive decline; Alzheimer's disease; bidirectional Mendelian randomization

## Subject area (primary)
Medicine — Epidemiology / Neurology / Pain medicine

## Author comments to readers (Research Square "Comments" field)
This preprint reports a bidirectional two-sample Mendelian randomization study of multisite chronic pain and cognitive outcomes. The China Health and Retirement Longitudinal Study (CHARLS) longitudinal validation is ongoing and flagged [PENDING DATA] throughout; those passages describe planned analyses, not results. Causal direction is left unresolved (significant reverse MR; non-significant Steiger test), and the headline cognitive-performance result is constrained by UK Biobank sample overlap between exposure and outcome; a non-UKB overlap-robustness re-interpretation is provided as corroborating evidence. Feedback on the methods and framing is welcome. A version 2 incorporating CHARLS data will be posted before journal submission (target: Scientific Reports, Nature Portfolio).

## Suggested journal for transfer (optional, set later)
Scientific Reports (Nature Portfolio) — to be selected at the transfer step after version 2 is posted.

## Files to upload (this package)
- Preprint_Manuscript.docx        (main text, with preprint banner)
- Supplementary_Information.docx  (S1 non-UKB sensitivity; S2 audit trail)
- STROBE-MR_checklist.docx        (reporting compliance)
- figures/Fig1_forest.pdf ... Fig5_nonukb.pdf  (or .png) — upload as separate figures
"""

def main():
    raw = open(B.SRC, encoding="utf-8").read()
    parts = raw.split("\n---\n", 1)
    body = parts[1] if len(parts) == 2 else raw
    body = B.renumber(body)
    body = re.sub(r"## Abstract\n.*?\n---\n", "## Abstract\n\n" + B.ABSTRACT + "\n\n---\n",
                  body, flags=re.S)
    bad = re.findall(r"[\u4e00-\u9fff]+", body)
    if bad:
        print("ABORT: Chinese in manuscript source:", sorted(set(bad))[:10])
        return 1
    paths = [build_preprint_manuscript(body)]
    # metadata + guide
    with open(os.path.join(OUT, "Preprint_Metadata.md"), "w", encoding="utf-8") as f:
        f.write(METADATA)
    paths.append(os.path.join(OUT, "Preprint_Metadata.md"))
    # copy SI, checklist, figures from the journal _upload bundle
    for fn in ("Supplementary_Information.docx", "STROBE-MR_checklist.docx"):
        src = os.path.join(SRC_UP, fn)
        if os.path.exists(src):
            shutil.copyfile(src, os.path.join(OUT, fn)); paths.append(os.path.join(OUT, fn))
    copied = 0
    for stem, (newstem, _d) in B.FIG_MAP.items():
        for ext in ("pdf", "png"):
            src = os.path.join(SRC_UP, "figures", f"{newstem}.{ext}")
            if os.path.exists(src):
                shutil.copyfile(src, os.path.join(OUTF, f"{newstem}.{ext}")); copied += 1
    # guide (written last so it can reference the built tree)
    guide = build_guide()
    paths.append(guide)
    print("built:")
    for p in sorted(paths):
        print(f"  {os.path.relpath(p, HERE):52s} {os.path.getsize(p)/1024:8.1f} KB")
    print(f"  figures copied: {copied}")
    return 0

def build_guide():
    g = """# 预印本上传指南（Research Square · 版本 1）

目标平台：**Research Square**（Springer Nature 旗下，与 *Scientific Reports* 同源，支持一键 transfer 到期刊；无编辑筛选、上传即获 DOI）。

> 注意：发预印本需用**你本人 ORCID** 登录 Research Square 操作，本指南仅把文件准备到位 + 给出分步指引。我无法代替你登录上传。

## 一、本包文件清单（`docs/submission/_upload_preprint/`）

| 文件 | 用途 | 上传位置 |
|---|---|---|
| `Preprint_Manuscript.docx` | 主稿件（含预印本横幅 + 版本说明） | Manuscript |
| `Supplementary_Information.docx` | S1 非 UKB 敏感性 + S2 审计轨 | Supplementary material |
| `STROBE-MR_checklist.docx` | STROBE-MR 2021 合规清单 | Supplementary material |
| `figures/Fig1_forest.pdf` … `Fig5_nonukb.pdf` | 5 张图（PDF 优先，PNG 备选） | Figures（逐张上传） |

## 二、分步上传

1. 打开 https://www.researchsquare.com ，用 **ORCID** 登录（首次需关联 ORCID 0009-0004-9698-6552）。
2. 点 **"Publish a preprint" / "Start new preprint"**。
3. **Manuscript**：上传 `Preprint_Manuscript.docx`。
4. **Figures**：逐张上传 `figures/` 下 5 个 PDF（或 PNG）。系统会自动编号 Fig 1–5。
5. **Supplementary material**：上传 `Supplementary_Information.docx` 与 `STROBE-MR_checklist.docx`。
6. **Title / Abstract / Authors / Keywords / Subject area**：照抄 `Preprint_Metadata.md` 对应字段。
   - Author：Yongxin Yang, **B.M.**（勿填 MD/PhD）；单位写全称。
   - Keywords 选 5–6 个。
7. **Comments to readers**：粘贴 `Preprint_Metadata.md` 中 "Author comments to readers" 段落（说明 CHARLS [PENDING DATA] 为前瞻性、因果方向未定、UKB 重叠限制）。
8. 确认 **License**（默认 CC BY 4.0，预印本推荐），提交。

## 三、版本与后续

- 上传后系统给 **DOI**，可引用为 "Yang Y. *Preprint.* Research Square. DOI:xxx"。
- **CHARLS 数据到位后**：用 `10_charls_extract.R` + `01_charls.R` 补全 §2.7/§3.5，升版本 A，在本稿基础上**再发版本 2**（Research Square 支持版本迭代，旧版本保留）。
- **版本 2 发布后**：在 Research Square 点 **"Transfer to journal"** → 选 *Scientific Reports*（或届时直接用 `docs/submission/_upload/` 的期刊投稿包走 Editorial Manager）。Springer Nature 期刊认可 Research Square 预印本，不视为既发。
- 数据可用性声明中的 `github.com/yyx-4113/pain-cognition-mr` 仓库请**先建好并含全部衍生数据 + 代码**，再上传预印本（否则声明失真）。

## 四、诚实性提醒（务必保留）

- 稿件中 `[PENDING DATA]` 段是**前瞻性计划**，不是结果；请勿在上传时改为"已完成"。
- 因果方向**未定**（反向 MR 显著、Steiger 不显著）；主认知结局受 UKB×UKB 重叠限制——这些限制已在正文与 Comments 中写明，上传前勿弱化。
"""
    path = os.path.join(OUT, "Preprint_Guide.md")
    with open(path, "w", encoding="utf-8") as f:
        f.write(g)
    return path

if __name__ == "__main__":
    raise SystemExit(main())
